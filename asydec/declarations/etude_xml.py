"""Étude descriptive persistante des XML, indépendante des simulations à huit articles."""
import datetime
import hashlib
import json
import math
import os
from pathlib import Path
from collections import defaultdict
import xml.etree.ElementTree as ET
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.formatting.rule import CellIsRule
from openpyxl.chart import BarChart, Reference

from asydec import chemins
from asydec.calcul import moteur

DOSSIER = chemins.ETUDES_XML


def analyser(chemin):
    contenu = Path(chemin).read_bytes()
    r = ET.fromstring(contenu)
    if r.tag != 'ASYCUDA':
        raise ValueError('Un XML ASYCUDA sans espace de noms est attendu.')
    identifiant = hashlib.sha256(contenu).hexdigest()
    alertes, lignes = [], []
    devise = moteur._txt(r, 'Valuation/Gs_Invoice/Currency_code') or 'NON RENSEIGNEE'
    def nombre(it, champ):
        t = moteur._txt(it, champ)
        if not t:
            return None
        valeur = float(t)
        if not math.isfinite(valeur):
            raise ValueError(f'Valeur non finie : {champ}')
        return valeur
    for position,it in enumerate(r.findall('Item'),1):
        nom = moteur._txt(it,'Goods_description/Description_of_goods')
        code = moteur._txt(it,'Tarification/HScode/Commodity_code')
        fob = nombre(it,'Tarification/Item_price')
        poids = nombre(it,'Valuation_item/Weight_itm/Gross_weight_itm')
        colis = nombre(it,'Packages/Number_of_packages')
        if not nom and not code and fob in (None,0) and poids in (None,0) and colis in (None,0):
            continue
        if poids is None or poids <= 0 or fob is None or fob <= 0:
            alertes.append(f'Art. {position} : poids ou FOB absent/non positif ; ratio exclu des statistiques.')
        if colis is not None and (colis < 0 or colis != int(colis)):
            alertes.append(f'Art. {position} : nombre de colis invalide.')
        lignes.append({'article':position,'code':code,'designation':nom,'poids':poids,'fob':fob,'colis':colis,
                       'ratio':fob/poids if poids is not None and poids>0 and fob is not None and fob>0 else None})
    if not lignes:
        raise ValueError('Aucun article exploitable dans ce XML.')
    if devise == 'NON RENSEIGNEE':
        alertes.append('Devise absente : comparaisons limitées aux fichiers de devise non renseignée.')
    for libelle,champ,cle,tolerance in [
        ('FOB','Valuation/Total/Total_invoice','fob',.005),
        ('Poids','Valuation/Total/Total_weight','poids',.0005),
        ('Colis','Property/Nbers/Total_number_of_packages','colis',0)]:
        total = nombre(r,champ)
        if total is not None and all(a[cle] is not None for a in lignes) and abs(sum(a[cle] for a in lignes)-total)>tolerance:
            alertes.append(f'{libelle} : somme des articles différente du total XML.')
    return {'id':identifiant,'fichier':Path(chemin).name,'source':str(Path(chemin).resolve()),
            'importe_utc':datetime.datetime.now(datetime.timezone.utc).isoformat(),
            'importateur':moteur._txt(r,'Traders/Consignee/Consignee_name').split('\n')[0],
            'reference':moteur._txt(r,'Declarant/Reference/Number'),
            'devise':devise,'articles':lignes,'alertes':alertes}


def enregistrer(chemin):
    etude = analyser(chemin)
    DOSSIER.mkdir(exist_ok=True)
    cible = DOSSIER / (etude['id']+'.json')
    if cible.exists():
        return json.loads(cible.read_text(encoding='utf-8'))
    # Copie de la source : les études restent consultables après déplacement de l'original.
    contenu = Path(chemin).read_bytes()
    if hashlib.sha256(contenu).hexdigest() != etude['id']:
        raise ValueError('Le fichier a changé pendant son analyse ; recommencez.')
    (DOSSIER / (etude['id']+'.xml')).write_bytes(contenu)
    temporaire = cible.with_suffix('.tmp')
    temporaire.write_text(json.dumps(etude,ensure_ascii=False,indent=2),encoding='utf-8')
    os.replace(temporaire,cible)
    return etude


def charger():
    etudes,erreurs = [],[]
    for p in sorted(DOSSIER.glob('*.json')) if DOSSIER.exists() else []:
        try:
            etudes.append(json.loads(p.read_text(encoding='utf-8')))
        except (ValueError,OSError) as exc:
            erreurs.append(f'{p.name} : {exc}')
    return etudes,erreurs


def groupe(etude,article):
    return (etude['devise'],article['code'], ' '.join(article['designation'].casefold().split()))


def comparer(etudes, reference_id, seuil=.10):
    """Compare à une référence explicite ; les hausses ne sont pas chronologiques."""
    if not math.isfinite(seuil) or seuil < 0:
        raise ValueError('Le seuil de variation doit être fini et positif ou nul.')
    reference = next((e for e in etudes if e['id']==reference_id),None)
    if reference is None:
        raise ValueError('Sélectionnez un XML de référence dans les études.')
    index=defaultdict(list)
    for a in reference['articles']:
        index[groupe(reference,a)].append(a)
    changements,problemes=[],[]
    for e in etudes:
        for alerte in e['alertes']:
            problemes.append({'fichier':e['fichier'],'article':'','niveau':'DONNÉE À VÉRIFIER','detail':alerte})
        if e['id']==reference_id:
            continue
        autres=defaultdict(list)
        for a in e['articles']:
            autres[groupe(e,a)].append(a)
        for cle in set(index)|set(autres):
            base,courant=index.get(cle,[]),autres.get(cle,[])
            if len(base)!=1 or len(courant)!=1:
                detail=('Correspondance ambiguë : code/désignation répétés.' if len(base)>1 or len(courant)>1 else
                        'Article absent du modèle comparé ou sans référence compatible (devise, code, désignation).')
                problemes.append({'fichier':e['fichier'],'article':cle[1]+' '+cle[2],'niveau':'NON COMPARABLE','detail':detail})
                continue
            a,b=base[0],courant[0]
            for metrique in ('poids','fob','colis','ratio'):
                avant,apres=a[metrique],b[metrique]
                delta=apres-avant if avant is not None and apres is not None else None
                relatif=delta/avant if delta is not None and avant>0 else None
                sens=('NON COMPARABLE' if delta is None else 'STABLE' if abs(delta)<1e-9 else 'HAUSSE' if delta>0 else 'BAISSE')
                important=relatif is not None and abs(relatif)>seuil+1e-12
                changements.append({'fichier':e['fichier'],'id':e['id'],'article':b['article'],'code':b['code'],
                    'designation':b['designation'],'devise':e['devise'],'metrique':metrique,'avant':avant,'apres':apres,
                    'delta':delta,'relatif':relatif,'sens':sens,'important':important})
                if important and metrique=='ratio':
                    problemes.append({'fichier':e['fichier'],'article':b['article'],'niveau':'VARIATION DE RATIO',
                        'detail':f"{sens} de {relatif:+.2%} par rapport à {reference['fichier']} (seuil {seuil:.1%}). À examiner, pas une erreur prouvée."})
    return {'reference':reference,'changements':changements,'problemes':problemes,'seuil':seuil}


def exporter(etudes,chemin,reference_id=None,seuil=.10):
    if not etudes:
        raise ValueError('Aucune étude à exporter.')
    etudes=list({e['id']:e for e in etudes}.values())
    reference_id = reference_id or etudes[0]['id']
    rapport = comparer(etudes,reference_id,seuil)
    wb = Workbook()
    aide = wb.active
    aide.title = 'AIDE'
    for ligne in [
        ['ANALYSE DES XML — données descriptives'],
        ['Source','XML importés ; nom, chemin et empreinte conservés dans MODELES.'],
        ['Ratio','FOB article / poids brut article. Unité : devise du XML par kg.'],
        ['Comparaison','Même devise, code marchandise et désignation normalisée. Aucun mélange de devises.'],
        ['Variations','Dispersion entre fichiers, sans interprétation chronologique.'],
        ['Écart relatif','(Maximum − minimum) / minimum ; vide si minimum nul.'],
        ['Moyenne pondérée','Somme FOB / somme poids sur les articles aux ratios exploitables.'],
        ['Prudence','Un XML étudié ne constitue pas une validation de conformité SYDONIA.'],
        ['Calculs','Formules Excel recalculées à l’ouverture. Les cellules vides signalent une donnée absente.'],
        ['Données','Une empreinte identique est comptée une seule fois. Les lignes totalement vides sont ignorées.']]:
        aide.append(ligne)
    aide.append(['Référence de comparaison',rapport['reference']['fichier']])
    aide.append(['Seuil de variation à examiner',seuil])
    aide['B12'].number_format='0.0%'
    aide.append(['Hausses et baisses','Écarts par rapport au XML choisi, sans ordre chronologique.'])
    aide.append(['Seuil','10 % par défaut, réglable par l’utilisateur ; indicateur descriptif, sans valeur réglementaire.'])
    aide.append(['Non comparable','Articles absents, doublons de code/désignation ou données manquantes : aucun rapprochement forcé.'])
    modeles=wb.create_sheet('MODELES')
    modeles.append(['ID SHA256','Fichier','Source','Importateur','Référence','Devise','Import UTC','Articles','Alertes'])
    articles=wb.create_sheet('ARTICLES')
    articles.append(['Modèle ID','Fichier','Importateur','Devise','Article XML','Code','Désignation','Poids kg','FOB','Colis','Ratio / kg'])
    alertes=wb.create_sheet('ALERTES')
    alertes.append(['Fichier','Alerte'])
    groupes=defaultdict(list)
    for e in etudes:
        modeles.append([e['id'],e['fichier'],e['source'],e['importateur'],e['reference'],e['devise'],e['importe_utc'],len(e['articles']),len(e['alertes'])])
        for alerte in e['alertes']:
            alertes.append([e['fichier'],alerte])
        for a in e['articles']:
            row=articles.max_row+1
            articles.append([e['id'],e['fichier'],e['importateur'],e['devise'],a['article'],a['code'],a['designation'],a['poids'],a['fob'],a['colis'],
                             f'=IF(AND(ISNUMBER(H{row}),ISNUMBER(I{row}),H{row}>0,I{row}>0),I{row}/H{row},"")'])
            if a['ratio'] is not None:
                groupes[groupe(e,a)].append(row)
    variations=wb.create_sheet('VARIATIONS')
    variations.append(['Devise','Code','Désignation normalisée','Observations','Ratio minimum','Ratio maximum','Ratio moyen','Ratio pondéré','Écart absolu','Écart relatif'])
    for cle,rows in sorted(groupes.items()):
        n=variations.max_row+1
        refs=','.join(f"'ARTICLES'!K{r}" for r in rows)
        fobs=','.join(f"'ARTICLES'!I{r}" for r in rows)
        poids=','.join(f"'ARTICLES'!H{r}" for r in rows)
        variations.append([*cle,f'=COUNT({refs})',f'=MIN({refs})',f'=MAX({refs})',f'=AVERAGE({refs})',
                           f'=SUM({fobs})/SUM({poids})',f'=F{n}-E{n}',f'=IF(E{n}>0,I{n}/E{n},"")'])
    changements=wb.create_sheet('CHANGEMENTS')
    changements.append(['Fichier','Article','Code','Désignation','Devise','Mesure','Référence','Valeur étudiée','Écart','Écart %','Sens','Seuil dépassé'])
    for c in rapport['changements']:
        n=changements.max_row+1
        changements.append([c['fichier'],c['article'],c['code'],c['designation'],c['devise'],c['metrique'],c['avant'],c['apres'],
            f'=IF(AND(ISNUMBER(G{n}),ISNUMBER(H{n})),H{n}-G{n},"")',
            f'=IF(AND(ISNUMBER(I{n}),G{n}>0),I{n}/G{n},"")',
            f'=IF(NOT(ISNUMBER(I{n})),"NON COMPARABLE",IF(ABS(I{n})<0.000000001,"STABLE",IF(I{n}>0,"HAUSSE","BAISSE")))',
            f'=IF(ISNUMBER(J{n}),IF(ABS(J{n})>AIDE!$B$12+0.000000000001,"À EXAMINER",""),"")'])
        changements.cell(n,10).number_format='0.00%'
    problemes=wb.create_sheet('PROBLEMES')
    problemes.append(['Fichier','Article','Signal','Explication'])
    for p in rapport['problemes']:
        problemes.append([p['fichier'],p['article'],p['niveau'],p['detail']])
    synthese=wb.create_sheet('SYNTHESE',0)
    synthese.append(['ANALYSE COMPLÈTE','Valeur'])
    synthese.append(['XML de référence',rapport['reference']['fichier']])
    synthese.append(['Modèles étudiés',len(etudes)])
    synthese.append(['Seuil de variation à examiner','=AIDE!B12'])
    synthese['B4'].number_format='0.0%'
    dernier=changements.max_row
    for libelle,critere in [('Hausses','HAUSSE'),('Baisses','BAISSE'),('Valeurs stables','STABLE')]:
        synthese.append([libelle,f'=COUNTIF(CHANGEMENTS!K2:K{max(2,dernier)},"{critere}")'])
    synthese.append(['Écarts au-dessus du seuil',f'=COUNTIF(CHANGEMENTS!L2:L{max(2,dernier)},"À EXAMINER")'])
    synthese.append(['Signaux à consulter',len(rapport['problemes'])])
    synthese.append(['Lecture','Consulter CHANGEMENTS pour les détails ; PROBLEMES pour les points à vérifier.'])
    synthese.append(['Limite','Les hausses sont relatives au XML choisi, sans notion de date. Un signal n’est pas une preuve d’erreur.'])
    if dernier>1:
        changements.conditional_formatting.add(f'J2:J{dernier}',CellIsRule(operator='greaterThan',formula=['0'],fill=PatternFill('solid',fgColor='FCE4D6')))
        changements.conditional_formatting.add(f'J2:J{dernier}',CellIsRule(operator='lessThan',formula=['0'],fill=PatternFill('solid',fgColor='DDEBF7')))
    # Les textes des XML restent des textes, y compris ceux commençant par '='.
    for ws in wb:
        ws.freeze_panes='A2'
        ws.auto_filter.ref=ws.dimensions
        for row in ws:
            for cell in row:
                if cell.data_type == 'f' and not ((ws.title=='ARTICLES' and cell.column==11) or (ws.title=='VARIATIONS' and cell.column>=4) or (ws.title=='CHANGEMENTS' and cell.column>=9) or (ws.title=='SYNTHESE' and cell.column==2 and cell.row in (4,5,6,7,8))):
                    cell.data_type='s'
                cell.font=Font(name='Arial',size=10)
        for c in ws[1]:
            c.font=Font(name='Arial',bold=True,color='FFFFFF',size=10)
            c.fill=PatternFill('solid',fgColor='1F4E79')
        for colonne in ws.columns:
            ws.column_dimensions[colonne[0].column_letter].width=min(65,max(16,max(len(str(c.value or '')) if c.data_type!='f' else 18 for c in colonne)+2))
    for row in articles.iter_rows(min_row=2,min_col=8,max_col=11):
        for c in row: c.number_format='0.000000'
    for row in variations.iter_rows(min_row=2,min_col=5,max_col=10):
        for c in row: c.number_format='0.000000'
        row[-1].number_format='0.00%'
    if variations.max_row>1:
        chart=BarChart()
        chart.title='Fourchettes des ratios par marchandise et devise'
        chart.add_data(Reference(variations,min_col=5,max_col=6,min_row=1,max_row=min(variations.max_row,21)),titles_from_data=True)
        chart.set_categories(Reference(variations,min_col=3,min_row=2,max_row=min(variations.max_row,21)))
        variations.add_chart(chart,'L2')
    wb.save(chemin)
    return str(chemin)
