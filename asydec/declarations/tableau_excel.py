"""Lecture d'un tableau d'articles (poids, FOB, colis) dans un fichier Excel quelconque."""
from openpyxl import load_workbook

from asydec.declarations import generer


def lire_tableau_excel(chemin):
    """Lit un tableau d'articles dans un fichier Excel : cherche une ligne d'en-tête
    contenant « Poids » (et si possible « FOB », « Colis »), puis les lignes d'articles
    jusqu'à « Total ». Renvoie {poids, fob, colis, noms} (listes de longueur variable)."""
    wb = load_workbook(chemin, data_only=True, read_only=True)
    try:
        for ws in wb.worksheets:
            lignes = list(ws.iter_rows(values_only=True))
            for h, ligne in enumerate(lignes):
                entetes = [str(c or "").strip().lower() for c in ligne]
                col = lambda mot: next((i for i, t in enumerate(entetes) if t.startswith(mot)), None)
                c_poids, c_fob, c_colis = col("poids"), col("fob"), col("colis")
                c_nom = col("marchandise")
                if c_poids is None:
                    continue
                res = {"poids": [], "fob": [], "colis": [], "noms": []}
                for suite in lignes[h + 1:]:
                    debut = " ".join(str(c or "") for c in suite[:2]).strip().lower()
                    if debut.startswith("total"):
                        break
                    p = generer.nombre(suite[c_poids]) if c_poids < len(suite) else None
                    if p is None:
                        continue
                    res["poids"].append(p)
                    res["fob"].append(generer.nombre(suite[c_fob]) if c_fob is not None else None)
                    res["colis"].append(generer.nombre(suite[c_colis]) if c_colis is not None else None)
                    res["noms"].append(str(suite[c_nom]) if c_nom is not None and suite[c_nom] else None)
                if res["poids"]:
                    return res
        raise ValueError("Aucun tableau avec une colonne « Poids » trouvé dans ce fichier.")
    finally:
        wb.close()
