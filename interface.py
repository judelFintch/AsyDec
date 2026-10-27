"""Interface graphique du générateur de déclarations ASYCUDA.

Lancer : double-clic sur « INTERFACE.command », ou
    .venv/bin/python interface.py
"""
import datetime
import os
import subprocess
import sys
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

import generer
import moteur
import variantes
import etude_xml

DOSSIER_SIMULATIONS = os.path.join(generer.ICI, "SIMULATIONS")


def lire_tableau_excel(chemin):
    """Lit un tableau d'articles dans un fichier Excel : cherche une ligne d'en-tête
    contenant « Poids » (et si possible « FOB », « Colis »), puis les lignes d'articles
    jusqu'à « Total ». Renvoie {poids, fob, colis, noms} (listes de 8 au plus)."""
    wb = load_workbook(chemin, data_only=True, read_only=True)
    try:
        for ws in wb.worksheets:
            lignes = list(ws.iter_rows(values_only=True))
            for h, ligne in enumerate(lignes):
                entetes = [str(c or "").strip().lower() for c in ligne]
                col = lambda mot: next((i for i, t in enumerate(entetes) if t.startswith(mot)), None)
                c_poids, c_fob, c_colis = col("poids"), col("fob"), col("colis")
                c_nom = col("marchandise")
                if c_poids is None:
                    continue
                res = {"poids": [], "fob": [], "colis": [], "noms": []}
                for suite in lignes[h + 1:]:
                    debut = " ".join(str(c or "") for c in suite[:2]).strip().lower()
                    if debut.startswith("total") or len(res["poids"]) == 8:
                        break
                    p = generer.nombre(suite[c_poids]) if c_poids < len(suite) else None
                    if p is None:
                        continue
                    res["poids"].append(p)
                    res["fob"].append(generer.nombre(suite[c_fob]) if c_fob is not None else None)
                    res["colis"].append(generer.nombre(suite[c_colis]) if c_colis is not None else None)
                    res["noms"].append(str(suite[c_nom]) if c_nom is not None and suite[c_nom] else None)
                if res["poids"]:
                    return res
        raise ValueError("Aucun tableau avec une colonne « Poids » trouvé dans ce fichier.")
    finally:
        wb.close()


class Application(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Déclarations ASYCUDA – DEC FRIPERIE")
        self.geometry("1120x860")
        self.minsize(1000, 760)
        self.ratios, self.table_imp = generer.charger_reglages()
        self.info = None
        self.dernier_fichier = None

        self.var_base = tk.StringVar()
        self.var_fob = tk.StringVar()
        self.var_colis = tk.StringVar()
        self.var_mode = tk.StringVar(value="RATIOS")
        self.var_dossier = tk.StringVar(value="Aucun fichier de base chargé.")
        self.var_totaux = tk.StringVar()
        self.var_sim_poids = tk.StringVar()
        self.var_sim_fob = tk.StringVar()
        self.var_sim_nb = tk.StringVar(value="Aucune simulation enregistrée.")
        self.simulations = []
        self.dernier_resultat = None
        self.controle_variantes = None
        self.var_controle = tk.StringVar()

        self._construire()

    # ------------------------------------------------------------------ UI
    def _construire(self):
        pad = {"padx": 8, "pady": 4}
        cadre = ttk.Frame(self, padding=12)
        cadre.pack(fill="both", expand=True)

        barre = ttk.Frame(cadre)
        barre.pack(fill="x", pady=(0, 6))
        ttk.Label(barre, text="Déclarations ASYCUDA", font=("TkDefaultFont", 12, "bold")).pack(side="left")
        ttk.Button(barre, text="Refresh — dernière version", command=self.rafraichir_projet).pack(side="right")

        # 1. Fichier de base
        f1 = ttk.LabelFrame(cadre, text="1. Fichier de base (déclaration SYDONIA du dossier)", padding=8)
        f1.pack(fill="x")
        ttk.Entry(f1, textvariable=self.var_base).pack(side="left", fill="x", expand=True, **pad)
        ttk.Button(f1, text="Parcourir…", command=self.parcourir).pack(side="left", **pad)
        ttk.Button(f1, text="Charger", command=self.charger_base).pack(side="left", **pad)
        ttk.Label(cadre, textvariable=self.var_dossier, foreground="#1F4E79").pack(fill="x", padx=8, pady=(4, 8))
        ttk.Label(cadre, textvariable=self.var_controle, foreground="#C55A11").pack(fill="x", padx=8)

        bibli = ttk.Frame(cadre)
        bibli.pack(fill="x")
        ttk.Button(bibli, text="Conserver des XML de contrôle…", command=self.importer_modeles).pack(side="left", padx=4)
        ttk.Button(bibli, text="Proposer des variantes…", command=self.proposer_variantes).pack(side="left", padx=4)

        ttk.Button(bibli, text="Étudier des XML…", command=self.etudier_xml).pack(side="left", padx=4)

        # 2. Paramètres
        f2 = ttk.LabelFrame(cadre, text="2. Paramètres", padding=8)
        f2.pack(fill="x")
        ttk.Label(f2, text="FOB total (USD)").grid(row=0, column=0, sticky="w", **pad)
        ttk.Entry(f2, textvariable=self.var_fob, width=14).grid(row=0, column=1, **pad)
        ttk.Label(f2, text="Colis total").grid(row=0, column=2, sticky="w", **pad)
        ttk.Entry(f2, textvariable=self.var_colis, width=10).grid(row=0, column=3, **pad)
        ttk.Label(f2, text="Mode FOB").grid(row=0, column=4, sticky="w", **pad)
        ttk.Radiobutton(f2, text="Ratios (poids × ratio, art. 4/7/8 ajustés)", value="RATIOS",
                        variable=self.var_mode, command=self._etat_fob).grid(row=0, column=5, sticky="w")
        ttk.Radiobutton(f2, text="FOB saisis", value="SAISI",
                        variable=self.var_mode, command=self._etat_fob).grid(row=0, column=6, sticky="w")

        # 3. Articles
        f3 = ttk.LabelFrame(cadre, text="3. Articles (poids du colisage)", padding=8)
        f3.pack(fill="x", pady=(8, 0))
        for c, t in enumerate(["Art.", "Marchandise", "Poids (kg)", "FOB (USD)", "Colis (vide = auto)", "Ratio"]):
            ttk.Label(f3, text=t, font=("TkDefaultFont", 11, "bold")).grid(row=0, column=c, sticky="w", padx=6)
        self.e_poids, self.e_fob, self.e_colis, self.l_ratio, self.l_nom = [], [], [], [], []
        for i in range(8):
            ttk.Label(f3, text=str(i + 1)).grid(row=i + 1, column=0, padx=6)
            l = ttk.Label(f3, text=generer.ARTICLES[i], width=28)
            l.grid(row=i + 1, column=1, sticky="w", padx=6)
            self.l_nom.append(l)
            for liste, col, w in ((self.e_poids, 2, 12), (self.e_fob, 3, 12), (self.e_colis, 4, 10)):
                e = ttk.Entry(f3, width=w, justify="right")
                e.grid(row=i + 1, column=col, padx=6, pady=2)
                liste.append(e)
            r = ttk.Label(f3, text=f"réf. {self.ratios[i]}", width=16)
            r.grid(row=i + 1, column=5, sticky="w", padx=6)
            self.l_ratio.append(r)
        ttk.Label(f3, textvariable=self.var_totaux, font=("TkDefaultFont", 11, "bold")).grid(
            row=9, column=1, columnspan=5, sticky="w", padx=6, pady=(6, 0))
        self._etat_fob()

        # Simulation
        fs = ttk.LabelFrame(cadre, text="Simulation (rien n'est enregistré tant que vous ne cliquez pas sur « Générer le XML »)", padding=8)
        fs.pack(fill="x", pady=(8, 0))
        ttk.Label(fs, text="Poids total voulu (kg)").pack(side="left", padx=4)
        ttk.Entry(fs, textvariable=self.var_sim_poids, width=10).pack(side="left", padx=4)
        ttk.Label(fs, text="FOB total voulu (USD)").pack(side="left", padx=4)
        ttk.Entry(fs, textvariable=self.var_sim_fob, width=10).pack(side="left", padx=4)
        ttk.Button(fs, text="Simuler", command=self.simuler).pack(side="left", padx=4)
        ttk.Button(fs, text="Ajouter à la comparaison", command=self.ajouter_simulation).pack(side="left", padx=4)
        ttk.Button(fs, text="Exporter les simulations (Excel)", command=self.exporter_simulations).pack(side="left", padx=4)
        ttk.Button(cadre, text="Utiliser le même poids, FOB et colis que le XML chargé", command=self.utiliser_totaux_base).pack(anchor="w", padx=8, pady=3)
        ttk.Label(cadre, textvariable=self.var_sim_nb, foreground="#7F6000").pack(fill="x", padx=8)

        # 4. Actions
        f4 = ttk.Frame(cadre, padding=(0, 10))
        f4.pack(fill="x")
        ttk.Button(f4, text="Importer un tableau Excel…", command=self.importer_excel).pack(side="left", padx=4)
        ttk.Button(f4, text="Reprendre les valeurs de la base", command=self.reprendre_base).pack(side="left", padx=4)
        ttk.Button(f4, text="Aperçu (sans enregistrer)", command=self.apercu).pack(side="left", padx=4)
        ttk.Button(f4, text="Générer le XML", command=self.generer).pack(side="left", padx=4)
        ttk.Button(f4, text="Afficher le fichier", command=self.afficher).pack(side="left", padx=4)
        ttk.Button(f4, text="Vider", command=self.vider).pack(side="right", padx=4)

        # 5. Résultat
        f5 = ttk.LabelFrame(cadre, text="Résultat", padding=8)
        f5.pack(fill="both", expand=True)
        self.txt = tk.Text(f5, height=12, wrap="word", font=("Menlo", 12))
        self.txt.pack(fill="both", expand=True)
        self._ecrire("Choisissez un fichier de base, saisissez les poids, puis « Aperçu » ou « Générer le XML ».")

    def _etat_fob(self):
        etat = "normal" if self.var_mode.get() == "SAISI" else "disabled"
        for e in self.e_fob:
            e.configure(state=etat)
        if etat == "disabled":
            self._fin_controle()

    def _fin_controle(self):
        """Quitte le contrôle d'une variante et revient aux ratios de SAISIE.xlsx."""
        if not self.controle_variantes:
            return
        self.controle_variantes = None
        self.ratios = generer.charger_reglages()[0]
        for k in range(8):
            self.l_ratio[k].configure(text=f"réf. {self.ratios[k]}")
        self.var_controle.set("")

    def _ecrire(self, texte):
        self.txt.configure(state="normal")
        self.txt.delete("1.0", "end")
        self.txt.insert("1.0", texte)
        self.txt.configure(state="disabled")

    @staticmethod
    def _val(e):
        t = e.get().strip()
        return generer.nombre(t) if t else None

    @staticmethod
    def _remplir(e, valeur):
        etat = str(e.cget("state"))
        e.configure(state="normal")
        e.delete(0, "end")
        if valeur is not None:
            e.insert(0, f"{valeur:.15g}" if isinstance(valeur, float) else str(valeur))
        e.configure(state=etat)

    # ------------------------------------------------------------- actions
    def parcourir(self):
        chemin = filedialog.askopenfilename(initialdir=generer.DEC, title="Fichier de base SYDONIA",
                                            filetypes=[("Déclarations XML", "*.xml"), ("Tous", "*.*")])
        if chemin:
            self.var_base.set(chemin)
            self.charger_base()

    def charger_base(self):
        try:
            chemin = generer.trouver_base(self.var_base.get())
            self.info = moteur.infos_dossier(moteur.charger(chemin).getroot())
        except Exception as e:  # noqa: BLE001
            messagebox.showerror("Fichier de base", str(e))
            return
        self._fin_controle()
        self.var_base.set(chemin)
        i = self.info
        self.var_dossier.set(
            f"{i['importateur']} ({i['code']})  •  réf. {i['reference']}  •  manifeste {i['manifeste']}  •  "
            f"camion {i['camion']}  •  taux {i['taux']:g}\n"
            f"Base : FOB {i['fob']:,.2f} USD  •  {i['poids']:,.0f} kg  •  {i['colis']} colis".replace(",", " "))
        for k, d in enumerate(i["designations"][:8]):
            self.l_nom[k].configure(text=(d or generer.ARTICLES[k])[:30])
        self.utiliser_totaux_base()

    def utiliser_totaux_base(self):
        """Reprend uniquement les totaux de la base, sans modifier les articles."""
        if not self.info:
            self.charger_base()
            return
        self.var_sim_poids.set(f"{self.info['poids']:.15g}")
        self.var_sim_fob.set(f"{self.info['fob']:.15g}")
        self.var_fob.set(f"{self.info['fob']:.15g}")
        self.var_colis.set(str(self.info['colis']))

    def reprendre_base(self):
        if not self.info:
            self.charger_base()
            if not self.info:
                return
        for k in range(8):
            self._remplir(self.e_poids[k], self.info["poids_articles"][k] if k < len(self.info["poids_articles"]) else None)
            self._remplir(self.e_fob[k], self.info["fob_articles"][k] if k < len(self.info["fob_articles"]) else None)
            self._remplir(self.e_colis[k], None)
        self.var_fob.set(f"{self.info['fob']:g}")
        self.var_colis.set(str(self.info["colis"]))

    def vider(self):
        for liste in (self.e_poids, self.e_fob, self.e_colis):
            for e in liste:
                self._remplir(e, None)
        self._fin_controle()
        self.var_totaux.set("")
        self._ecrire("Champs vidés.")

    def _calculer(self, enregistrer):
        if not self.info:
            self.charger_base()
            if not self.info:
                return None
        try:
            colis = [self._val(e) for e in self.e_colis]
            if self.controle_variantes:
                if self.var_mode.get() != "SAISI":
                    raise ValueError("La variante contrôlée doit rester en mode FOB saisis.")
                if variantes.signature(self.var_base.get()) != self.controle_variantes['signature']:
                    raise ValueError("Les articles de la base diffèrent du contrôle actif.")
                valeurs_controle = {'poids':[self._val(e) for e in self.e_poids],
                                    'fob':[self._val(e) for e in self.e_fob], 'colis':colis}
                variantes.controler(valeurs_controle, self.controle_variantes)
                if (abs(sum(valeurs_controle['fob'])-generer.nombre(self.var_fob.get())) > .005 or
                        sum(colis) != generer.nombre(self.var_colis.get())):
                    raise ValueError("Les totaux saisis ne correspondent plus à la variante.")
            res = generer.produire(
                self.var_base.get(),
                [self._val(e) for e in self.e_poids],
                self.ratios, self.table_imp, self.registre_charge(),
                mode=self.var_mode.get(),
                fob_total=generer.nombre(self.var_fob.get()) if self.var_fob.get().strip() else None,
                fob_saisis=[self._val(e) for e in self.e_fob],
                colis_total=generer.nombre(self.var_colis.get()) if self.var_colis.get().strip() else None,
                colis_saisis=colis if all(c is not None for c in colis) else None,
                enregistrer=enregistrer,
            )
        except Exception as e:  # noqa: BLE001
            messagebox.showerror("Calcul impossible", str(e))
            return None
        if enregistrer:
            generer.ecrire_registre(self._registre)
        self._afficher_resultat(res, enregistrer)
        self.dernier_resultat = res
        return res

    # ------------------------------------------------------- import Excel
    def importer_excel(self):
        chemin = filedialog.askopenfilename(initialdir=os.path.expanduser("~/Downloads"),
                                            title="Tableau Excel (poids, FOB, colis)",
                                            filetypes=[("Excel", "*.xlsx *.xlsm"), ("Tous", "*.*")])
        if not chemin:
            return
        try:
            t = lire_tableau_excel(chemin)
        except Exception as e:  # noqa: BLE001
            messagebox.showerror("Import Excel", str(e))
            return
        n = len(t["poids"])
        for k in range(8):
            self._remplir(self.e_poids[k], t["poids"][k] if k < n else None)
            self._remplir(self.e_fob[k], t["fob"][k] if k < n else None)
            self._remplir(self.e_colis[k], int(t["colis"][k]) if k < n and t["colis"][k] is not None else None)
        fob_complets = n == 8 and all(f is not None for f in t["fob"])
        if fob_complets:
            self.var_mode.set("SAISI")
            self.var_fob.set(f"{round(sum(t['fob']), 2):g}")
        self._etat_fob()
        if n == 8 and all(c is not None for c in t["colis"]):
            self.var_colis.set(str(int(sum(t["colis"]))))
        self._ecrire(f"Tableau importé : {os.path.basename(chemin)}\n"
                     f"{n} articles • {sum(t['poids']):,.0f} kg".replace(",", " ")
                     + (f" • FOB {sum(t['fob']):,.2f} USD (mode « FOB saisis »)".replace(",", " ") if fob_complets else "")
                     + ("\n⚠ Le tableau n'a pas 8 articles." if n != 8 else "")
                     + "\n\nVérifiez les valeurs, choisissez le fichier de base, puis « Aperçu » ou « Simuler ».")

    # ---------------------------------------------------------- simulation
    def simuler(self):
        poids = [self._val(e) for e in self.e_poids]
        if any(p is None for p in poids):
            messagebox.showerror("Simulation", "Saisissez ou importez d'abord les 8 poids.")
            return
        cible_p = generer.nombre(self.var_sim_poids.get()) if self.var_sim_poids.get().strip() else None
        cible_f = generer.nombre(self.var_sim_fob.get()) if self.var_sim_fob.get().strip() else None
        if cible_p is None and cible_f is None:
            messagebox.showerror("Simulation", "Indiquez un poids total et/ou un FOB total voulus.")
            return
        if cible_p is not None:
            poids = moteur.ajuster_poids(poids, cible_p)
            for k in range(8):
                self._remplir(self.e_poids[k], poids[k])
        if cible_f is not None:
            self.var_mode.set("RATIOS")
            self.var_fob.set(f"{cible_f:g}")
            self._etat_fob()
        res = self._calculer(enregistrer=False)
        if res:
            self.txt.configure(state="normal")
            self.txt.insert("1.0", "SIMULATION – poids ramenés au même pourcentage, FOB selon les ratios.\n")
            self.txt.configure(state="disabled")

    def ajouter_simulation(self):
        res = self._calculer(enregistrer=False)
        if not res:
            return
        r = res["resume"]
        self.simulations.append({
            "nom": f"Simulation {len(self.simulations) + 1}",
            "dossier": f"{res['info']['importateur']} – {res['info']['reference']}",
            "poids": [a["poids"] for a in r["articles"]], "fob": res["fob"], "colis": res["colis"],
            "cif": [a["cif"] for a in r["articles"]], "taxes": [a["taxes"] for a in r["articles"]],
            "taxes_globales": r["taxes_globales"], "taxes_totales": r["taxes_totales"],
            "alertes": res["alertes"],
        })
        self.var_sim_nb.set(f"{len(self.simulations)} simulation(s) dans la comparaison.")

    def exporter_simulations(self):
        if not self.simulations:
            messagebox.showinfo("Simulations", "Ajoutez d'abord au moins une simulation à la comparaison.")
            return
        os.makedirs(DOSSIER_SIMULATIONS, exist_ok=True)
        chemin = os.path.join(DOSSIER_SIMULATIONS,
                              f"Simulations {datetime.datetime.now():%Y-%m-%d %H%M}.xlsx")
        wb = Workbook()
        ws = wb.active
        ws.title = "COMPARAISON"
        ws.append(["SIMULATIONS – document de travail, ce n'est pas une déclaration"])
        ws["A1"].font = Font(bold=True, size=13, color="C00000")
        ws.append([f"Créé le {datetime.datetime.now():%d/%m/%Y %H:%M}"])
        ws.append([])
        entete = ["Art.", "Marchandise"]
        for s in self.simulations:
            entete += [f"{s['nom']} – Poids", "FOB", "Ratio", "Colis", "Taxes CDF"]
        ws.append(entete)
        jaune = PatternFill("solid", fgColor="FFF2CC")
        for c in ws[4]:
            c.font = Font(bold=True)
            c.fill = jaune
        for k in range(8):
            ligne = [k + 1, generer.ARTICLES[k]]
            for s in self.simulations:
                p, f = s["poids"][k], s["fob"][k]
                ligne += [p, f, round(f / p, 4) if p else None, s["colis"][k], s["taxes"][k]]
            ws.append(ligne)
        tot = ["", "TOTAL"]
        for s in self.simulations:
            tot += [sum(s["poids"]), round(sum(s["fob"]), 2), None, sum(s["colis"]), sum(s["taxes"])]
        ws.append(tot)
        glob = ["", "Taxes globales"]
        for s in self.simulations:
            glob += ["", "", "", "", s["taxes_globales"]]
        ws.append(glob)
        tt = ["", "TOTAL TAXES"]
        for s in self.simulations:
            tt += ["", "", "", "", s["taxes_totales"]]
        ws.append(tt)
        for c in ws[ws.max_row]:
            c.font = Font(bold=True)
        ws.append([])
        for s in self.simulations:
            ws.append([s["nom"], s["dossier"], "Alertes : " + (" | ".join(s["alertes"]) or "aucune")])
        ws.column_dimensions["B"].width = 26
        wb.save(chemin)
        subprocess.run(["open", chemin], check=False)
        self.var_sim_nb.set(f"{len(self.simulations)} simulation(s) exportée(s) : {os.path.basename(chemin)}")

    def etudier_xml(self):
        fen = tk.Toplevel(self)
        fen.title("Étude des XML — ratios et variations")
        fen.geometry("1050x720")
        statut = tk.StringVar()
        ttk.Label(fen, textvariable=statut).pack(fill="x", padx=10, pady=5)
        tableau = ttk.Treeview(fen, columns=("fichier","importateur","articles","devise","alertes"), show="headings", height=9)
        for cle,titre in [("fichier","Fichier"),("importateur","Importateur"),("articles","Articles"),("devise","Devise"),("alertes","Alertes")]:
            tableau.heading(cle,text=titre)
            tableau.column(cle,width=320 if cle=="fichier" else 150)
        tableau.pack(fill="both",expand=True,padx=10)
        detail=tk.Text(fen,height=17,wrap="word",font=("Menlo",11))
        detail.pack(fill="both",expand=True,padx=10,pady=8)
        etudes=[]

        def afficher(event=None):
            if not tableau.selection():
                return
            indice=tableau.selection()[0]
            if not indice.isdigit() or int(indice)>=len(etudes):
                return
            e=etudes[int(indice)]
            detail.delete("1.0","end")
            detail.insert("end",f"{e['fichier']}\nSource : {e['source']}\nDevise : {e['devise']} — ratio = FOB / poids brut (kg)\n\n")
            for a in e['articles']:
                ratio=f"{a['ratio']:.8g}" if a['ratio'] is not None else "non calculable"
                detail.insert("end",f"Art. {a['article']} — {a['code']} — {a['designation']}\nPoids : {a['poids']} kg | FOB : {a['fob']} | Colis : {a['colis']} | Ratio : {ratio}\n\n")
            detail.insert("end","Alertes : " + (" | ".join(e['alertes']) or "aucune"))

        def actualiser():
            donnees,erreurs=etude_xml.charger()
            etudes[:]=donnees
            tableau.delete(*tableau.get_children())
            for k,e in enumerate(etudes):
                tableau.insert("","end",iid=str(k),values=(e['fichier'],e['importateur'],len(e['articles']),e['devise'],len(e['alertes'])))
            statut.set(f"{len(etudes)} XML conservé(s). Sélectionnez une ligne pour consulter ses ratios. Export : toute la bibliothèque d’études.")
            if etudes:
                tableau.selection_set("0")
                afficher()
            if erreurs:
                messagebox.showwarning("Études illisibles","\n".join(erreurs),parent=fen)

        def importer(chemins):
            erreurs=[]
            for chemin in chemins:
                try:
                    etude_xml.enregistrer(chemin)
                except Exception as exc:
                    erreurs.append(f"{os.path.basename(chemin)} : {exc}")
            actualiser()
            if erreurs:
                messagebox.showwarning("Fichiers non étudiés","\n".join(erreurs),parent=fen)

        def fichiers():
            importer(filedialog.askopenfilenames(parent=fen,title="XML à étudier (sélection multiple)",filetypes=[("XML","*.xml")]))

        def dossier():
            from pathlib import Path
            chemin=filedialog.askdirectory(parent=fen,title="Dossier à étudier, sous-dossiers inclus")
            if chemin:
                importer([p for p in Path(chemin).rglob('*') if p.is_file() and p.suffix.lower()=='.xml'])

        def reference_et_seuil():
            if not tableau.selection():
                raise ValueError("Sélectionnez le XML qui servira de référence.")
            seuil=generer.nombre(var_seuil.get())
            if seuil is None or seuil<0:
                raise ValueError("Indiquez un seuil de variation positif ou nul, en %.")
            return etudes[int(tableau.selection()[0])]['id'],seuil/100

        def analyser_complet():
            try:
                reference,seuil=reference_et_seuil()
                rapport=etude_xml.comparer(etudes,reference,seuil)
                detail.delete("1.0","end")
                detail.insert("end",f"ANALYSE COMPLÈTE — référence : {rapport['reference']['fichier']}\nSeuil : {seuil:.1%}. Écarts par rapport au modèle choisi, sans chronologie.\n")
                for sens in ('HAUSSE','BAISSE','STABLE','NON COMPARABLE'):
                    detail.insert("end",f"{sens} : {sum(c['sens']==sens for c in rapport['changements'])} mesures\n")
                detail.insert("end","\nPrincipaux écarts de ratios (plus forts en premier) :\n")
                ratios=sorted((c for c in rapport['changements'] if c['metrique']=='ratio' and c['relatif'] is not None),key=lambda c:abs(c['relatif']),reverse=True)
                for c in ratios[:30]:
                    detail.insert("end",f"{c['fichier']} — art. {c['article']} : {c['sens']} {c['relatif']:+.2%} ({c['avant']:.6g} → {c['apres']:.6g})" + (" — À EXAMINER" if c['important'] else "") + "\n")
                detail.insert("end","\nPoints à vérifier :\n")
                for p in rapport['problemes']:
                    detail.insert("end",f"{p['fichier']} — {p['niveau']} : {p['detail']}\n")
                if not rapport['problemes']:
                    detail.insert("end","Aucun signal détecté par ces contrôles.\n")
            except Exception as exc:
                messagebox.showerror("Analyse",str(exc),parent=fen)

        def exporter():
            if not etudes:
                messagebox.showinfo("Analyse","Importez d’abord des XML à étudier.",parent=fen)
                return
            chemin=filedialog.asksaveasfilename(parent=fen,title="Exporter les données et variations",defaultextension=".xlsx",
                initialfile=f"Analyse XML {datetime.datetime.now():%Y-%m-%d %H%M%S}.xlsx",filetypes=[("Excel","*.xlsx")])
            if chemin:
                try:
                    reference,seuil=reference_et_seuil()
                    etude_xml.exporter(etudes,chemin,reference,seuil)
                    messagebox.showinfo("Export terminé",f"Analyse enregistrée :\n{chemin}\nLes formules se recalculent à l’ouverture dans Excel.",parent=fen)
                except Exception as exc:
                    messagebox.showerror("Export impossible",str(exc),parent=fen)

        reglages=ttk.Frame(fen)
        reglages.pack(fill="x",padx=10)
        ttk.Label(reglages,text="Référence = XML sélectionné dans la liste. Seuil à examiner (%) :").pack(side="left")
        var_seuil=tk.StringVar(value="10")
        ttk.Entry(reglages,textvariable=var_seuil,width=8).pack(side="left",padx=4)
        ttk.Button(reglages,text="Analyse complète",command=analyser_complet).pack(side="left",padx=8)
        actions=ttk.Frame(fen)
        actions.pack(fill="x",padx=10,pady=8)
        ttk.Button(actions,text="Ajouter des XML…",command=fichiers).pack(side="left",padx=4)
        ttk.Button(actions,text="Ajouter un dossier…",command=dossier).pack(side="left",padx=4)
        ttk.Button(actions,text="Exporter l’analyse Excel…",command=exporter).pack(side="left",padx=4)
        tableau.bind("<<TreeviewSelect>>",afficher)
        actualiser()

    def importer_modeles(self):
        chemins = filedialog.askopenfilenames(title="XML de référence", filetypes=[("XML", "*.xml")])
        erreurs, total = [], 0
        for chemin in chemins:
            try:
                variantes.conserver(chemin)
                total += 1
            except Exception as exc:
                erreurs.append(f"{os.path.basename(chemin)} : {exc}")
        if chemins:
            messagebox.showinfo("Bibliothèque", f"{total} modèle(s) conservé(s).\n" + "\n".join(erreurs))

    def proposer_variantes(self):
        if not self.info:
            self.charger_base()
            if not self.info:
                return
        fen = tk.Toplevel(self)
        fen.title("Variantes contrôlées — classement par proximité des références")
        fen.geometry("900x650")
        disponibles = variantes.modeles()
        ttk.Label(fen, text="XML de contrôle conservé (même ordre et mêmes marchandises que la base)").pack()
        choix = ttk.Combobox(fen, state="readonly", values=[p.name for p in disponibles], width=85)
        choix.pack(fill="x", padx=10)
        if disponibles:
            choix.current(0)
        champs = []
        for titre, valeur in [("Poids total (kg)", f"{self.info['poids']:.15g}"),
                              ("FOB total (USD)", f"{self.info['fob']:.15g}"),
                              ("Colis total", str(self.info['colis']))]:
            ligne = ttk.Frame(fen)
            ligne.pack(fill="x", padx=10, pady=3)
            ttk.Label(ligne, text=titre, width=22).pack(side="left")
            entree = ttk.Entry(ligne)
            entree.insert(0, valeur)
            entree.pack(side="left")
            champs.append(entree)
        def reprendre_totaux():
            for entree, cle in zip(champs, ('poids', 'fob', 'colis')):
                self._remplir(entree, self.info[cle])
        ttk.Button(fen, text="Utiliser le même poids, FOB et colis que le XML chargé", command=reprendre_totaux).pack(anchor="w", padx=10, pady=4)
        ttk.Label(fen, text="Art. 4/7/8 : ratios du XML choisi. Autres : fourchettes des XML compatibles. Arrondis au centime.").pack()
        tableau = ttk.Treeview(fen, columns=("score", "poids", "fob", "colis", "taxes"), show="headings", height=8)
        for cle, titre in zip(("score", "poids", "fob", "colis", "taxes"), ("Proximité / 100", "Poids kg", "FOB USD", "Colis", "Taxes estimées CDF")):
            tableau.heading(cle, text=titre)
            tableau.column(cle, width=140)
        tableau.pack(fill="x", padx=10, pady=8)
        detail = tk.Text(fen, height=10, wrap="word")
        detail.pack(fill="both", expand=True, padx=10)
        candidats = []
        contexte = {}
        deja_proposes = {}

        def selection(event=None):
            if not tableau.selection() or tableau.selection()[0] not in tableau.get_children():
                return
            v = candidats[int(tableau.selection()[0])]
            detail.delete("1.0", "end")
            profil = contexte['profil']
            detail.insert("end", f"Proximité : {v['proximite']}/100 — référence : {v['reference_proche']}\n")
            detail.insert("end", "Score = 100 × (1 − distance moyenne des parts de poids, FOB et colis).\n")
            detail.insert("end", f"{len(profil['references'])} contrôle(s) compatible(s). Classement parmi les propositions trouvées.\n")
            detail.insert("end", "Art.    Poids kg    FOB USD    Colis    Ratio obtenu    Règle de contrôle\n")
            for k in range(8):
                bas,haut = profil['bornes'][k]
                regle = f"fixe {bas:.6g}" if k in moteur.ARTICLES_VARIABLES else f"[{bas:.6g} ; {haut:.6g}]"
                detail.insert("end", f"{k+1}    {v['poids'][k]:.3f}    {v['fob'][k]:.2f}    {v['colis'][k]}    {v['fob'][k]/v['poids'][k]:.6g}    {regle}\n")
            if len(profil['references']) == 1:
                detail.insert("end", "Une seule référence : les fourchettes sont réduites à un ratio unique.\n")
            if profil['exclus']:
                detail.insert("end", "Contrôles exclus : " + " | ".join(profil['exclus']) + "\n")
            detail.insert("end", "\nAlertes : " + (" | ".join(v['resultat']['alertes']) or "aucune"))

        def calculer(regenerer=False):
            try:
                if choix.current() < 0:
                    raise ValueError("Conservez d'abord un XML de contrôle.")
                controle = str(disponibles[choix.current()])
                info = variantes.verifier_modele(controle)
                base = generer.trouver_base(self.var_base.get())
                variantes.verifier_modele(base)
                if variantes.signature(base) != variantes.signature(controle):
                    raise ValueError("La base et le contrôle doivent avoir les mêmes codes et désignations dans le même ordre.")
                profil = variantes.construire_profil(controle)
                ratios = profil['ratios']
                valeurs = [generer.nombre(e.get()) for e in champs]
                cle = (base, controle, tuple(valeurs), tuple(profil['bornes']))
                historique = deja_proposes.get(cle, set())
                propositions = variantes.generer(*valeurs, ratios, profil=profil,
                    nombre=1 if regenerer else 12, exclure=historique)

                registre = generer.lire_registre()
                for v in propositions:
                    ratios_v = [ratios[k] if k in moteur.ARTICLES_VARIABLES else v['fob'][k]/v['poids'][k] for k in range(8)]
                    v['resultat'] = generer.produire(base, v['poids'], ratios_v, self.table_imp, registre,
                        mode="SAISI", fob_total=valeurs[1], fob_saisis=v['fob'],
                        colis_total=valeurs[2], colis_saisis=v['colis'], enregistrer=False)
                deja_proposes.setdefault(cle, set()).update(variantes.cle_variante(v) for v in propositions)
                candidats[:] = propositions
                contexte.update(base=base, ratios=ratios, profil=profil, controle=os.path.basename(controle))
                tableau.delete(*tableau.get_children())
                for k,v in enumerate(candidats):
                    tableau.insert("", "end", iid=str(k), values=(v["proximite"], f"{sum(v['poids']):.3f}",
                        f"{sum(v['fob']):.2f}", sum(v['colis']), v['resultat']['resume']['taxes_totales']))
                tableau.selection_set("0")
                selection()
            except Exception as exc:
                messagebox.showerror("Variantes", str(exc), parent=fen)

        def appliquer():
            if not tableau.selection() or tableau.selection()[0] not in tableau.get_children():
                return
            v = candidats[int(tableau.selection()[0])]
            self.var_base.set(contexte['base'])
            self.charger_base()
            self.controle_variantes = contexte['profil']
            self.ratios = [contexte['ratios'][k] if k in moteur.ARTICLES_VARIABLES else v['fob'][k]/v['poids'][k] for k in range(8)]
            self.var_controle.set(f"CONTRÔLE ACTIF ({contexte['controle']}) : ratios de la variante. "
                                  "« Vider », « Charger » ou le mode Ratios pour en sortir.")
            self.var_mode.set("SAISI")
            self._etat_fob()
            for k in range(8):
                for entrees, cle in [(self.e_poids,"poids"), (self.e_fob,"fob"), (self.e_colis,"colis")]:
                    self._remplir(entrees[k], v[cle][k])
            self.var_fob.set(f"{sum(v['fob']):.2f}")
            self.var_colis.set(str(sum(v['colis'])))
            self._calculer(False)
            fen.destroy()

        tableau.bind("<<TreeviewSelect>>", selection)
        actions = ttk.Frame(fen)
        actions.pack(fill="x", pady=8)
        ttk.Button(actions, text="Créer 12 variantes", command=calculer).pack(side="left", padx=10)
        ttk.Button(actions, text="Régénérer un autre modèle", command=lambda: calculer(True)).pack(side="left", padx=10)
        ttk.Button(actions, text="Reprendre la variante choisie", command=appliquer).pack(side="left", padx=10)

    def registre_charge(self):
        self._registre = generer.lire_registre()
        return self._registre

    def _afficher_resultat(self, res, enregistre):
        r = res["resume"]
        for k in range(8):
            if self.var_mode.get() != "SAISI":
                self._remplir(self.e_fob[k], res["fob"][k])
            p = r["articles"][k]["poids"]
            self.l_ratio[k].configure(text=f"{res['fob'][k] / p:.4f}  (réf. {self.ratios[k]})" if p else "")
        self.var_totaux.set(f"Total : {r['poids']:,.0f} kg   •   {r['fob']:,.2f} USD   •   "
                            f"{sum(res['colis'])} colis".replace(",", " "))
        lignes = ["ENREGISTRÉ : " + os.path.relpath(res["chemin"], generer.DEC) if enregistre
                  else "APERÇU (rien n'est enregistré) – fichier qui serait créé : " + res["nom"], ""]
        if self.controle_variantes:
            score = variantes.evaluer({'poids':[a['poids'] for a in r['articles']],
                                      'fob':res['fob'], 'colis':res['colis']}, self.controle_variantes)
            lignes.append(f"CONTRÔLE ACTIF : ratios 4/7/8 fixes, autres dans les fourchettes (arrondis au centime).")
            lignes.append(f"Proximité : {score['proximite']}/100 — référence : {score['reference_proche']}")
        lignes.append(f"{'Art':<4}{'Poids':>9}{'FOB':>11}{'Colis':>7}{'CIF CDF':>14}{'Taxes CDF':>13}")
        for k, a in enumerate(r["articles"]):
            lignes.append(f"{k+1:<4}{a['poids']:>9,.0f}{a['fob']:>11,.2f}{res['colis'][k]:>7}"
                          f"{a['cif']:>14,}{a['taxes']:>13,}".replace(",", " "))
        lignes += ["",
                   f"Taxes articles : {r['taxes_articles']:,} CDF   +   taxes globales : {r['taxes_globales']:,} CDF".replace(",", " "),
                   f"TOTAL TAXES : {r['taxes_totales']:,} CDF".replace(",", " "), ""]
        lignes += ["Alertes :"] + [f"  ⚠ {a}" for a in res["alertes"]] if res["alertes"] else ["Alertes : aucune ✔"]
        self._ecrire("\n".join(lignes))

    def apercu(self):
        self._calculer(enregistrer=False)

    def generer(self):
        res = self._calculer(enregistrer=False)
        if not res:
            return
        message = f"Créer le fichier :\n\n{res['nom']}\n\ndans {os.path.dirname(os.path.relpath(res['chemin'], generer.DEC))} ?"
        if res["alertes"]:
            message += "\n\nAlertes :\n" + "\n".join(f"• {a}" for a in res["alertes"])
        if not messagebox.askyesno("Générer le XML", message):
            return
        res = self._calculer(enregistrer=True)
        if res:
            self.dernier_fichier = res["chemin"]
            messagebox.showinfo("Fichier créé", res["nom"])

    def rafraichir_projet(self):
        """Relance l'interpréteur pour recharger tous les modules depuis le disque."""
        if not messagebox.askyesno(
            "Recharger la dernière version",
            "Recharger la version du projet présente sur cet ordinateur ?\n\n"
            "Les saisies et simulations non enregistrées seront perdues. "
            "Les XML, modèles et études déjà enregistrés seront conservés.",
            parent=self,
        ):
            return
        try:
            os.execv(sys.executable, [sys.executable, os.path.abspath(__file__)])
        except OSError as exc:
            messagebox.showerror("Rechargement impossible", str(exc), parent=self)

    def afficher(self):
        if self.dernier_fichier and os.path.exists(self.dernier_fichier):
            subprocess.run(["open", "-R", self.dernier_fichier], check=False)
        else:
            subprocess.run(["open", generer.DEC], check=False)


if __name__ == "__main__":
    Application().mainloop()
