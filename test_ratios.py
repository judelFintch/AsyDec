import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook
import ratios


def xml(articles, importateur='IMPORT A'):
    items = ''.join(f'<Item><Tarification><Item_price>{fob}</Item_price><HScode><Commodity_code>{code}</Commodity_code></HScode></Tarification>'
                    f'<Goods_description><Description_of_goods>{nom}</Description_of_goods></Goods_description>'
                    f'<Valuation_item><Weight_itm><Gross_weight_itm>{poids}</Gross_weight_itm></Weight_itm></Valuation_item></Item>'
                    for code, nom, poids, fob in articles)
    return (f'<ASYCUDA><Traders><Consignee><Consignee_name>{importateur}</Consignee_name></Consignee></Traders>'
            f'<Valuation><Gs_Invoice><Currency_code>USD</Currency_code></Gs_Invoice></Valuation>{items}<Item/></ASYCUDA>')


class RatiosTest(unittest.TestCase):
    def fichiers(self, dossier):
        contenus = [
            xml([('6309', 'Friperie', 100, 150), ('6304', 'Couvre-lits', 200, 40)]),
            xml([('6304', '  COUVRE-LITS ', 100, 15), ('6309', 'Friperie', 100, 140)], 'IMPORT B'),
            xml([('6309', 'Friperie', 0, 0)]),
            '<autre/>']
        chemins = []
        for k, contenu in enumerate(contenus):
            p = Path(dossier) / f'{k}.xml'
            p.write_text(contenu)
            chemins.append(p)
        doublon = Path(dossier) / 'copie.xml'
        doublon.write_text(contenus[0])
        return chemins + [doublon]

    def test_regroupe_par_article_quelle_que_soit_la_position(self):
        with tempfile.TemporaryDirectory() as d:
            observations, erreurs = ratios.extraire(self.fichiers(d))
            self.assertEqual(len(observations), 4)  # copie identique et article sans poids ignorés
            self.assertEqual(len(erreurs), 1)       # fichier qui n'est pas un XML ASYCUDA
            groupes = ratios.regrouper(observations)
            # Positions à égalité (1 et 2) : ordre alphabétique.
            self.assertEqual([g['designation'] for g in groupes], ['Couvre-lits', 'Friperie'])
            couvre, friperie = groupes
            self.assertEqual(len(couvre['observations']), 2)  # casse et espaces ignorés
            self.assertEqual((friperie['min'], friperie['max']), (1.4, 1.5))
            self.assertAlmostEqual(friperie['pondere'], 290 / 200)
            self.assertEqual((couvre['min'], couvre['max'], couvre['mediane']), (.15, .2, .175))
            self.assertAlmostEqual(couvre['pondere'], 55 / 300)
            self.assertAlmostEqual(couvre['ecart'], 1 / 3)

    def test_separation_par_importateur_et_export(self):
        with tempfile.TemporaryDirectory() as d:
            observations, erreurs = ratios.extraire(self.fichiers(d))
            groupes = ratios.regrouper(observations, par_importateur=True)
            self.assertEqual(len(groupes), 4)
            self.assertEqual({g['importateur'] for g in groupes}, {'IMPORT A', 'IMPORT B'})
            sortie = Path(d) / 'ratios.xlsx'
            ratios.exporter(groupes, sortie, erreurs)
            wb = load_workbook(sortie)
            self.assertEqual(wb.sheetnames, ['RATIOS', 'DETAIL', 'FICHIERS ECARTES'])
            self.assertEqual(wb['RATIOS'].max_row, 5)
            self.assertEqual(wb['DETAIL'].max_row, 5)


if __name__ == '__main__':
    unittest.main()
