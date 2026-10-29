import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import load_workbook

from asydec.declarations import etude_xml

XML = '''<ASYCUDA><Valuation><Gs_Invoice><Currency_code>USD</Currency_code></Gs_Invoice><Total><Total_invoice>150</Total_invoice><Total_weight>100</Total_weight></Total></Valuation>
<Item><Tarification><Item_price>150</Item_price><HScode><Commodity_code>6309</Commodity_code></HScode></Tarification><Goods_description><Description_of_goods>Friperie</Description_of_goods></Goods_description><Valuation_item><Weight_itm><Gross_weight_itm>100</Gross_weight_itm></Weight_itm></Valuation_item><Packages><Number_of_packages>10</Number_of_packages></Packages></Item><Item/>
<Item><Tarification><Item_price>0</Item_price></Tarification><Goods_description><Description_of_goods>Article sans poids</Description_of_goods></Goods_description></Item></ASYCUDA>'''


class EtudeTest(unittest.TestCase):
    def test_persistance_dedoublonnage_et_export(self):
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)/'base.xml'
            p.write_text(XML)
            with patch.object(etude_xml,'DOSSIER',Path(tmp)/'etudes'):
                a=etude_xml.enregistrer(p)
                etude_xml.enregistrer(p)
                p.unlink()
                donnees,erreurs=etude_xml.charger()
                self.assertEqual(len(donnees),1)
                self.assertEqual(erreurs,[])
                self.assertEqual(len(a['articles']),2)
                self.assertEqual(a['articles'][0]['ratio'],1.5)
                self.assertIsNone(a['articles'][1]['ratio'])
                self.assertTrue(a['alertes'])
                self.assertTrue((etude_xml.DOSSIER/(a['id']+'.xml')).exists())
                sortie=Path(tmp)/'analyse.xlsx'
                etude_xml.exporter(donnees,sortie)
                wb=load_workbook(sortie)
                self.assertEqual(wb.sheetnames,['SYNTHESE','AIDE','MODELES','ARTICLES','ALERTES','VARIATIONS','CHANGEMENTS','PROBLEMES'])
                self.assertEqual(wb['ARTICLES']['H2'].value,100)
                self.assertEqual(wb['ARTICLES']['I2'].value,150)
                self.assertIn('I2/H2',wb['ARTICLES']['K2'].value)
                self.assertEqual(wb['VARIATIONS']['E2'].value,"=MIN('ARTICLES'!K2)")
                self.assertEqual(wb['VARIATIONS'].max_row,2)
                wb.close()

    def test_refuse_xml_inattendu(self):
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)/'autre.xml'
            p.write_text('<autre/>')
            with self.assertRaises(ValueError):
                etude_xml.analyser(p)

    def test_groupes_par_devise_et_marchandise(self):
        article={'code':'6309','designation':'  AUTRES   Articles '}
        self.assertEqual(etude_xml.groupe({'devise':'USD'},article),('USD','6309','autres articles'))
        self.assertNotEqual(etude_xml.groupe({'devise':'USD'},article),etude_xml.groupe({'devise':'EUR'},article))



class ComparaisonTest(unittest.TestCase):
    def etude(self,identifiant,ratio,devise='USD'):
        return {'id':identifiant,'fichier':identifiant+'.xml','devise':devise,'alertes':[],
            'articles':[{'article':1,'code':'6309','designation':'Friperie','poids':100,'fob':ratio*100,'ratio':ratio,'colis':10}]}

    def test_hausse_baisse_et_seuil(self):
        etudes=[self.etude('ref',1),self.etude('haut',1.3),self.etude('bas',.8)]
        r=etude_xml.comparer(etudes,'ref',.1)
        ratios=[c for c in r['changements'] if c['metrique']=='ratio']
        self.assertEqual([c['sens'] for c in ratios],['HAUSSE','BAISSE'])
        self.assertTrue(all(c['important'] for c in ratios))
        self.assertEqual(len(r['problemes']),2)
        self.assertAlmostEqual(ratios[0]['relatif'],.3)

    def test_devise_et_doublons_non_comparables(self):
        ref=self.etude('ref',1)
        autre=self.etude('autre',2,'EUR')
        r=etude_xml.comparer([ref,autre],'ref')
        self.assertEqual(r['changements'],[])
        autre=self.etude('autre',2)
        autre['articles']*=2
        r=etude_xml.comparer([ref,autre],'ref')
        self.assertEqual(r['changements'],[])
        self.assertTrue(r['problemes'])

    def test_reference_et_seuil_invalides(self):
        with self.assertRaises(ValueError):
            etude_xml.comparer([self.etude('ref',1)],'absent')
        with self.assertRaises(ValueError):
            etude_xml.comparer([self.etude('ref',1)],'ref',float('nan'))


if __name__=='__main__':
    unittest.main()
