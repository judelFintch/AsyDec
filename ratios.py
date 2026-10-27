"""Récupération des ratios : lit un groupe de XML et regroupe les ratios FOB/kg par article.

Un article est reconnu par sa devise, son code marchandise et sa désignation
normalisée (casse et espaces ignorés), quelle que soit sa position dans le XML.
Un même fichier présent plusieurs fois (même contenu) n'est compté qu'une fois.
"""
import statistics
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

import etude_xml


def extraire(chemins):
    """Renvoie (observations, erreurs) : une observation par article valorisé de chaque XML."""
    observations, erreurs, vus = [], [], set()
    for chemin in chemins:
        try:
            etude = etude_xml.analyser(chemin)
        except Exception as exc:  # noqa: BLE001 - un fichier illisible n'arrête pas les autres
            erreurs.append(f"{Path(chemin).name} : {exc}")
            continue
        if etude["id"] in vus:
            continue
        vus.add(etude["id"])
        for a in etude["articles"]:
            if a["ratio"] is None:
                continue
            observations.append({
                "fichier": etude["fichier"], "chemin": str(chemin), "importateur": etude["importateur"],
                "devise": etude["devise"], "position": a["article"], "code": a["code"],
                "designation": a["designation"], "poids": a["poids"], "fob": a["fob"], "ratio": a["ratio"],
                "groupe": etude_xml.groupe(etude, a)})
    return observations, erreurs


def regrouper(observations, par_importateur=False):
    """Statistiques par article, dans l'ordre habituel des articles (position la plus fréquente)."""
    groupes = defaultdict(list)
    for o in observations:
        cle = o["groupe"] + ((o["importateur"],) if par_importateur else ())
        groupes[cle].append(o)
    resultats = []
    for cle, obs in groupes.items():
        ratios = [o["ratio"] for o in obs]
        bas, haut = min(ratios), max(ratios)
        positions = [o["position"] for o in obs]
        resultats.append({
            "devise": obs[0]["devise"], "code": obs[0]["code"], "designation": obs[0]["designation"],
            "importateur": obs[0]["importateur"] if par_importateur else "",
            "position": min(set(positions), key=lambda n: (-positions.count(n), n)),  # égalité : la plus petite
            "observations": sorted(obs, key=lambda o: o["ratio"]),
            "fichiers": len({o["fichier"] for o in obs}),
            "min": bas, "max": haut, "moyenne": statistics.fmean(ratios), "mediane": statistics.median(ratios),
            "pondere": sum(o["fob"] for o in obs) / sum(o["poids"] for o in obs),
            "ecart": (haut - bas) / bas if bas > 0 else None})
    return sorted(resultats, key=lambda g: (g["importateur"], g["position"], g["designation"].casefold()))


def exporter(groupes, chemin, erreurs=()):
    """Classeur Excel : synthèse par article, détail par fichier, fichiers écartés."""
    wb = Workbook()
    entete = Font(bold=True, color="FFFFFF")
    fond = PatternFill("solid", fgColor="1F4E79")
    synthese = wb.active
    synthese.title = "RATIOS"
    synthese.append(["Art.", "Code", "Désignation", "Importateur", "Devise", "Observations", "Fichiers",
                     "Ratio min", "Ratio max", "Moyenne", "Médiane", "Pondéré (ΣFOB/Σkg)", "Écart max/min"])
    detail = wb.create_sheet("DETAIL")
    detail.append(["Désignation", "Code", "Fichier", "Importateur", "Art. dans le XML", "Poids kg", "FOB", "Ratio"])
    for g in groupes:
        synthese.append([g["position"], g["code"], g["designation"], g["importateur"], g["devise"],
                         len(g["observations"]), g["fichiers"], g["min"], g["max"], g["moyenne"],
                         g["mediane"], g["pondere"], g["ecart"]])
        for o in g["observations"]:
            detail.append([g["designation"], o["code"], o["fichier"], o["importateur"], o["position"],
                           o["poids"], o["fob"], o["ratio"]])
    if erreurs:
        ecartes = wb.create_sheet("FICHIERS ECARTES")
        ecartes.append(["Fichier et raison"])
        for e in erreurs:
            ecartes.append([e])
    for ws in wb:
        for cellule in ws[1]:
            cellule.font, cellule.fill = entete, fond
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for colonne in ws.columns:
            largeur = max(len(str(c.value)) if c.value is not None else 0 for c in colonne)
            ws.column_dimensions[colonne[0].column_letter].width = min(60, max(10, largeur + 2))
    for ligne in synthese.iter_rows(min_row=2, min_col=8, max_col=12):
        for c in ligne:
            c.number_format = "0.0000"
    for c in synthese["M"][1:]:
        c.number_format = "0.0%"
    for c in detail["H"][1:]:
        c.number_format = "0.0000"
    wb.save(chemin)
    return str(chemin)
