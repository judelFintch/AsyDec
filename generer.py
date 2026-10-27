"""Génère les déclarations XML ASYCUDA à partir du fichier Excel SAISIE.xlsx.

Utilisation : double-cliquer sur « LANCER.command », ou
    .venv/bin/python generer.py

Chaque ligne de la feuille SAISIE dont le statut est « A GENERER » produit
un fichier XML, rangé dans DEC FRIPERIE/<IMPORTATEUR>/1 - HISTORIQUE/.
Le résultat (nom du fichier, taxes, alertes) est réécrit dans la ligne.
"""
import datetime
import json
import os
import re
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

import moteur
import variantes

ICI = os.path.dirname(os.path.abspath(__file__))
DEC = os.path.dirname(ICI)
EXCEL = os.path.join(ICI, "SAISIE.xlsx")
REGISTRE = os.path.join(ICI, "registre.json")

ARTICLES = ["Friperie", "Couvre-lits", "Vêtements F/F", "Moustiquaires",
            "Couvertures", "Parties vêtements", "Articles pêche", "Articles ménage"]
RATIOS_DEFAUT = [1.5, 0.1875, 0.25, 0.400, 0.15, 0.136, 0.419, 0.157]
IMPORTATEURS_DEFAUT = [
    ("SOCIETE K.D.META", "KDMETA"),
    ("SOCIETE QUATRE E", "QUATRE E"),
    ("SHABANI SUMAILI CHARLY", "SHABANI SUMAILI"),
    ("HUSSEIN ZEIN", "HUSSEIN ZEIN"),
]
SEUIL_RATIO = 0.10  # alerte si un ratio variable s'écarte de plus de 10 %

# Colonnes de la feuille SAISIE
COL = {"statut": 1, "base": 2, "fob": 3, "colis_total": 4, "mode": 5,
       "poids": 6, "fob_art": 14, "colis_art": 22,
       "fichier": 30, "fob_final": 31, "poids_final": 32, "taxes": 33, "alertes": 34, "date": 35}


# --------------------------------------------------------------------------
# Modèle Excel
# --------------------------------------------------------------------------
def creer_modele():
    wb = Workbook()
    ws = wb.active
    ws.title = "SAISIE"
    entetes = (["Statut", "Fichier de base (nom ou chemin)", "FOB total (USD)", "Colis total", "Mode FOB"]
               + [f"Poids {i+1} {a}" for i, a in enumerate(ARTICLES)]
               + [f"FOB {i+1} {a}" for i, a in enumerate(ARTICLES)]
               + [f"Colis {i+1} {a}" for i, a in enumerate(ARTICLES)]
               + ["→ Fichier généré", "→ FOB final", "→ Poids final", "→ Taxes totales (CDF)", "→ Alertes", "→ Date"])
    ws.append(entetes)
    bleu, gris = PatternFill("solid", fgColor="DDEBF7"), PatternFill("solid", fgColor="EDEDED")
    for c in range(1, len(entetes) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.fill = gris if c >= COL["fichier"] else bleu
    ws.row_dimensions[1].height = 45
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 42
    ws.column_dimensions["AD"].width = 40
    ws.column_dimensions["AH"].width = 60
    ws.freeze_panes = "C2"
    ws.append(["EXEMPLE", "shabani 19000KG 18.xml", 7800, 790, "RATIOS",
               550, 280, 280, 7100, 90, 490, 7050, 3160])
    dv = DataValidation(type="list", formula1='"A GENERER,FAIT,EXEMPLE,IGNORER"', allow_blank=True)
    dv2 = DataValidation(type="list", formula1='"RATIOS,SAISI"', allow_blank=True)
    ws.add_data_validation(dv)
    ws.add_data_validation(dv2)
    dv.add("A2:A500")
    dv2.add("E2:E500")

    wr = wb.create_sheet("RATIOS")
    wr.append(["Article", "Marchandise", "Ratio FOB/kg", "Type"])
    for i, (a, r) in enumerate(zip(ARTICLES, RATIOS_DEFAUT)):
        wr.append([i + 1, a, r, "variable" if i in moteur.ARTICLES_VARIABLES else "fixe"])
    for c in "ABCD":
        wr[f"{c}1"].font = Font(bold=True)
    wr.column_dimensions["B"].width = 22

    wi = wb.create_sheet("IMPORTATEURS")
    wi.append(["Nom dans le XML (destinataire)", "Nom du dossier"])
    for nom, dossier in IMPORTATEURS_DEFAUT:
        wi.append([nom, dossier])
    wi["A1"].font = wi["B1"].font = Font(bold=True)
    wi.column_dimensions["A"].width = 36
    wi.column_dimensions["B"].width = 24

    wa = wb.create_sheet("AIDE")
    for ligne in [
        "MODE D'EMPLOI",
        "1. Feuille SAISIE : une ligne par déclaration.",
        "2. Statut = A GENERER pour les lignes à traiter (EXEMPLE et FAIT sont ignorées).",
        "3. Fichier de base : le XML SYDONIA du dossier (nom seul : il est cherché dans DEC FRIPERIE).",
        "4. Poids 1 à 8 : les poids du colisage, dans l'ordre des articles du fichier de base.",
        "5. Mode RATIOS : FOB calculés dans les bornes des XML ; ratios 4/7/8 ajustables, autres dans les fourchettes.",
        "   Mode SAISI : les FOB 1 à 8 sont contrôlés avec les mêmes règles ; incohérence = refus.",
        "6. Colis : saisir les 8 entiers positifs, ou aucun : estimation graduée (minimum 10, tous différents).",
        "   (vide = total du fichier de base).",
        "7. Lancer : double-clic sur LANCER.command. Les colonnes grises se remplissent.",
        "8. Le fichier est rangé dans DEC FRIPERIE/<IMPORTATEUR>/1 - HISTORIQUE, nommé",
        "   IMPORTATEUR FOBUSD POIDSKG NN.xml (NN = numéro suivant). Rien n'est jamais écrasé.",
    ]:
        wa.append([ligne])
    wa["A1"].font = Font(bold=True, size=13)
    wa.column_dimensions["A"].width = 110
    wb.save(EXCEL)


# --------------------------------------------------------------------------
# Utilitaires
# --------------------------------------------------------------------------
def trouver_base(nom):
    if not nom:
        raise ValueError("Fichier de base manquant.")
    nom = str(nom).strip().strip('"')
    if os.path.isabs(nom) and os.path.exists(nom):
        return nom
    trouves = []
    for dossier, _, fichiers in os.walk(DEC):
        if os.sep + "OUTILS" in dossier + os.sep and dossier.startswith(ICI):
            continue
        if nom in fichiers:
            trouves.append(os.path.join(dossier, nom))
    if not trouves:
        raise ValueError(f"Fichier de base introuvable dans DEC FRIPERIE : {nom}")
    # Priorité aux fichiers de base de l'utilisateur
    trouves.sort(key=lambda p: (0 if "VOS FICHIERS" in p else 1, p))
    return trouves[0]


def nom_dossier(importateur, table):
    for nom, dossier in table:
        if nom and importateur.upper().startswith(str(nom).upper()):
            return dossier
    return re.sub(r"^SOCIETE\s+", "", importateur.upper()).replace(".", "").strip() or "INCONNU"


# Rangement d'un importateur dans DEC FRIPERIE/<IMPORTATEUR>/ : un sous-dossier par type de fichier.
RANGEMENT = {"historique": "1 - HISTORIQUE", "sources": "2 - VOS FICHIERS (bases et SYDONIA)",
             "valides": "4 - XML VALIDES", "pdf": "5 - PDF"}  # « 3 - A ESSAYER » existe déjà chez certains


def dossier_rangement(importateur, cle, table_imp=None, creer=True):
    """Sous-dossier de rangement de l'importateur ; créé au besoin, jamais vidé ni déplacé."""
    table = IMPORTATEURS_DEFAUT if table_imp is None else table_imp
    chemin = os.path.join(DEC, nom_dossier(importateur or "", table), RANGEMENT[cle])
    if creer:
        os.makedirs(chemin, exist_ok=True)
    return chemin


def numero_suivant(dossier_imp):
    motif = re.compile(re.escape(dossier_imp) + r" .*KG (\d{2})(?:-\d{2})?(?: - .*)?\.xml$", re.I)
    maxi = 0
    for _, _, fichiers in os.walk(DEC):
        for f in fichiers:
            m = motif.match(f)
            if m:
                maxi = max(maxi, int(m.group(1)))
    return maxi + 1


def fmt_nom(x):
    return str(int(x)) if abs(x - round(x)) < 1e-9 else f"{x:.2f}"


def nombre(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    return float(str(v).replace(" ", "").replace(",", "."))


def lire_registre():
    try:
        with open(REGISTRE, encoding="utf-8") as f:
            return json.load(f)
    except (OSError, ValueError):
        return {}


def ecrire_registre(reg):
    with open(REGISTRE, "w", encoding="utf-8") as f:
        json.dump(reg, f, ensure_ascii=False, indent=1)


# --------------------------------------------------------------------------
# Calcul d'une déclaration (utilisé par Excel et par l'interface)
# --------------------------------------------------------------------------
def produire(base, poids, ratios, table_imp, registre, mode="RATIOS", fob_total=None,
             fob_saisis=None, colis_total=None, colis_saisis=None, enregistrer=True, profil=None, plaque=None):
    """Calcule une déclaration. Si enregistrer=False, rien n'est écrit (aperçu).
    Renvoie un dictionnaire : chemin, nom, resume, alertes, fob, colis, info."""
    alertes = []
    arbre = moteur.charger(base)
    info = moteur.infos_dossier(arbre.getroot())
    plaque = (plaque or "").strip() or None
    if plaque:
        # Plaque saisie à l'étape Dossier : elle remplace celle du XML de base.
        ancienne = info["camion"]
        plaque = moteur.ecrire_plaque(arbre.getroot(), plaque)
        info = moteur.infos_dossier(arbre.getroot())
        if plaque != ancienne:
            alertes.append(f"Plaque inscrite dans le XML : {plaque} (XML de base : {ancienne or 'aucune'}).")
    n = len(info["fob_articles"])
    if n != 8:
        alertes.append(f"Le fichier de base a {n} articles (attendu 8).")
    poids = list(poids)[:n]
    if len(poids) < n or any(p is None for p in poids):
        raise ValueError("Poids manquant pour un article.")
    mode = (mode or "RATIOS").strip().upper()

    profil = profil or variantes.construire_profil(base)
    if variantes.signature(base) != profil['signature']:
        raise ValueError('Les articles de la base sont incompatibles avec le contrôle choisi.')
    if mode == "SAISI":
        fob = list(fob_saisis or [])[:n]
        if len(fob)!=8 or any(f is None for f in fob):
            raise ValueError('Un FOB saisi est requis pour chacun des huit articles.')
    elif mode == "RATIOS":
        fob = variantes.fob_controles(poids,profil,fob_total)
    else:
        raise ValueError('Mode inconnu : choisir RATIOS ou SAISI.')
    if fob_total is not None and abs(sum(fob)-fob_total)>.005:
        raise ValueError('La somme des FOB ne correspond pas au total demandé.')
    saisis = list(colis_saisis or [])
    if any(c is not None for c in saisis):
        if len(saisis)!=8 or any(c is None for c in saisis):
            raise ValueError('Renseignez les huit colis ou laissez les huit champs vides.')
        colis=saisis
    else:
        total_colis = info['colis'] if colis_total is None else colis_total
        if int(total_colis)!=total_colis:
            raise ValueError('Le total des colis doit être entier.')
        colis=moteur.colis_gradues(poids,fob,int(total_colis))
        alertes.append('Colis estimés au poids : minimum 10 par article, tous différents. À vérifier avec le colisage réel.')
    variantes.controler({'poids':poids,'fob':fob,'colis':colis},profil)
    if colis_total is not None and sum(colis)!=colis_total:
        raise ValueError('La somme des colis ne correspond pas au total demandé.')
    if profil.get('variables_libres'):
        for k in moteur.ARTICLES_VARIABLES:
            obtenu=fob[k]/poids[k]
            ecart=obtenu/profil['ratios'][k]-1
            if abs(ecart)>.10:
                alertes.append(f'Art. {k+1} : ratio ajusté {obtenu:.6g}, référence {profil["ratios"][k]:.6g}, écart {ecart:+.1%}. À examiner (seuil indicatif 10 %).')


    if abs(sum(fob) - info["fob"]) > 0.005:
        alertes.append(f"FOB total {sum(fob):.2f} ≠ base {info['fob']:.2f}.")
    if abs(sum(poids) - info["poids"]) > 0.0005:
        alertes.append(f"Poids total {sum(poids):g} ≠ base {info['poids']:g}.")

    # Sans camion ni manifeste, impossible de reconnaître le dossier : pas de suivi.
    cle = f"{info['camion']} | {info['manifeste']}" if info["camion"] or info["manifeste"] else None
    precedent = registre.get(cle) if cle else None
    if precedent and [round(p, 3) for p in precedent["poids"]] != [round(p, 3) for p in poids]:
        alertes.append(f"ATTENTION : poids par article différents du dernier fichier généré pour ce camion "
                       f"({precedent['fichier']}). Vérifier le colisage.")

    resume = moteur.recalculer(arbre, list(zip(fob, poids)), colis)
    dossier_imp = nom_dossier(info["importateur"], table_imp)
    dest = os.path.join(DEC, dossier_imp, RANGEMENT["historique"])
    nn = numero_suivant(dossier_imp)
    while True:
        nom = f"{dossier_imp} {fmt_nom(resume['fob'])}USD {fmt_nom(resume['poids'])}KG {nn:02d}.xml"
        chemin = os.path.join(dest, nom)
        if not os.path.exists(chemin):
            break
        nn += 1
    if enregistrer:
        os.makedirs(dest, exist_ok=True)
        os.makedirs(os.path.join(DEC, dossier_imp, RANGEMENT["sources"]), exist_ok=True)
        moteur.enregistrer(arbre, chemin)
        if cle:
            registre[cle] ={"fichier": nom, "poids": poids, "date": datetime.datetime.now().isoformat(timespec="minutes")}
    return {"chemin": chemin, "nom": nom, "resume": resume, "alertes": alertes,
            "fob": fob, "colis": colis, "info": info, "mode": mode, "plaque": plaque}


# --------------------------------------------------------------------------
# Traitement d'une ligne Excel
# --------------------------------------------------------------------------
def traiter(ws, row, ratios, table_imp, registre):
    v = lambda c: ws.cell(row=row, column=c).value
    res = produire(
        trouver_base(v(COL["base"])),
        [nombre(v(COL["poids"] + i)) for i in range(8)],
        ratios, table_imp, registre,
        mode=str(v(COL["mode"]) or "RATIOS"),
        fob_total=nombre(v(COL["fob"])),
        fob_saisis=[nombre(v(COL["fob_art"] + i)) for i in range(8)],
        colis_total=nombre(v(COL["colis_total"])),
        colis_saisis=[nombre(v(COL["colis_art"] + i)) for i in range(8)],
    )
    resume, alertes = res["resume"], res["alertes"]
    ws.cell(row=row, column=COL["statut"]).value = "FAIT"
    ws.cell(row=row, column=COL["fichier"]).value = os.path.relpath(res["chemin"], DEC)
    ws.cell(row=row, column=COL["fob_final"]).value = resume["fob"]
    ws.cell(row=row, column=COL["poids_final"]).value = resume["poids"]
    ws.cell(row=row, column=COL["taxes"]).value = resume["taxes_totales"]
    ws.cell(row=row, column=COL["alertes"]).value = " | ".join(alertes) or "Aucune"
    ws.cell(row=row, column=COL["date"]).value = datetime.datetime.now().strftime("%d/%m/%Y %H:%M")
    for i, c in enumerate(res["colis"]):
        ws.cell(row=row, column=COL["colis_art"] + i).value = c
    if res["mode"] != "SAISI":
        for i, f in enumerate(res["fob"]):
            ws.cell(row=row, column=COL["fob_art"] + i).value = f
    return res["nom"], resume, alertes


def charger_reglages():
    """Ratios et table des importateurs depuis SAISIE.xlsx (valeurs par défaut sinon)."""
    ratios, table_imp = list(RATIOS_DEFAUT), list(IMPORTATEURS_DEFAUT)
    if os.path.exists(EXCEL):
        try:
            wb = load_workbook(EXCEL, read_only=True)
            ratios = [float(wb["RATIOS"].cell(row=i + 2, column=3).value) for i in range(8)]
            table_imp = [(r[0], r[1]) for r in wb["IMPORTATEURS"].iter_rows(min_row=2, values_only=True) if r[0]]
            wb.close()
        except Exception:  # noqa: BLE001 - réglages par défaut si le fichier est illisible
            pass
    return ratios, table_imp


def main():
    if not os.path.exists(EXCEL):
        creer_modele()
        print(f"Modèle créé : {EXCEL}\nRemplissez la feuille SAISIE puis relancez.")
        return 0
    try:
        wb = load_workbook(EXCEL)
    except PermissionError:
        print("Fermez SAISIE.xlsx dans Excel puis relancez.")
        return 1
    ws = wb["SAISIE"]
    ratios = [float(wb["RATIOS"].cell(row=i + 2, column=3).value) for i in range(8)]
    table_imp = [(r[0], r[1]) for r in wb["IMPORTATEURS"].iter_rows(min_row=2, values_only=True) if r[0]]
    registre = lire_registre()

    faits = erreurs = 0
    for row in range(2, ws.max_row + 1):
        statut = str(ws.cell(row=row, column=COL["statut"]).value or "").strip().upper()
        if statut != "A GENERER":
            continue
        try:
            nom, resume, alertes = traiter(ws, row, ratios, table_imp, registre)
            faits += 1
            print(f"✔ Ligne {row} : {nom}  |  taxes {resume['taxes_totales']:,} CDF".replace(",", " "))
            for a in alertes:
                print(f"    ⚠ {a}")
        except Exception as e:  # noqa: BLE001 - on veut continuer les autres lignes
            erreurs += 1
            ws.cell(row=row, column=COL["statut"]).value = "ERREUR"
            ws.cell(row=row, column=COL["alertes"]).value = str(e)
            print(f"✘ Ligne {row} : {e}")

    try:
        wb.save(EXCEL)
    except PermissionError:
        print("⚠ Impossible d'écrire SAISIE.xlsx (ouvert dans Excel ?). Les XML sont quand même générés.")
    ecrire_registre(registre)
    if faits == 0 and erreurs == 0:
        print("Aucune ligne « A GENERER » dans la feuille SAISIE.")
    else:
        print(f"\nTerminé : {faits} fichier(s) généré(s), {erreurs} erreur(s).")
    return 1 if erreurs else 0


if __name__ == "__main__":
    sys.exit(main())
